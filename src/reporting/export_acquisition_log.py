"""
export_acquisition_log.py

Generate a formatted Excel view of the authoritative SQLite acquisition log.
SQLite remains the system of record. The Excel workbook is a reproducible
report enriched with metadata from the resolved Source Registry.
"""

from pathlib import Path
from tempfile import NamedTemporaryFile
import sqlite3

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.config import ACQUISITION_LOG_FILE, DATABASE_FILE
from src.database.schema import ACQUISITION_LOG_TABLE, initialize_database_schema
from src.governance.source_registry import build_resolved_registry
from src.utils.logger import logger

SHEET_NAME = "Acquisition Log"
TABLE_NAME = "AcquisitionLogTable"
JOIN_STATUS_COLUMN = "Registry Match Status"

EXPORT_COLUMNS = [
    "Run ID",
    "Execution Date UTC",
    "Process Type",
    "Dataset ID",
    "Source ID",
    "Dataset Name",
    "Production Method",
    "Output File",
    "Storage Location",
    "Status",
    "Records Produced",
    JOIN_STATUS_COLUMN,
    "Notes",
    "Created At UTC",
]

HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUCCESS_FILL = PatternFill(fill_type="solid", fgColor="C6EFCE")
PARTIAL_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")
FAILED_FILL = PatternFill(fill_type="solid", fgColor="F4CCCC")
PLANNED_FILL = PatternFill(fill_type="solid", fgColor="D9EAF7")
UNRESOLVED_FILL = PatternFill(fill_type="solid", fgColor="FCE4D6")
THIN_GRAY_SIDE = Side(style="thin", color="D9E1F2")

COLUMN_WIDTHS = {
    "Run ID": 31,
    "Execution Date UTC": 22,
    "Process Type": 16,
    "Dataset ID": 13,
    "Source ID": 12,
    "Dataset Name": 28,
    "Production Method": 19,
    "Output File": 30,
    "Storage Location": 42,
    "Status": 17,
    "Records Produced": 18,
    JOIN_STATUS_COLUMN: 21,
    "Notes": 60,
    "Created At UTC": 22,
}


def read_acquisition_log():
    """Read the authoritative acquisition log from SQLite."""
    initialize_database_schema()
    connection = sqlite3.connect(DATABASE_FILE)
    try:
        query = f"""
            SELECT
                run_id,
                execution_date,
                process_type,
                output_file,
                storage_location,
                status,
                records_produced,
                notes,
                created_at
            FROM {ACQUISITION_LOG_TABLE}
            ORDER BY execution_date DESC, run_id DESC
        """
        return pd.read_sql_query(query, connection)
    except sqlite3.DatabaseError as error:
        raise RuntimeError("Unable to read the SQLite Acquisition Log.") from error
    finally:
        connection.close()


def build_registry_view():
    """Build the unique registry metadata view used for enrichment."""
    registry = build_resolved_registry()
    required_columns = [
        "Output File Name",
        "Dataset ID",
        "Source ID",
        "Dataset Name",
        "Production Method",
    ]
    missing_columns = [
        column for column in required_columns if column not in registry.columns
    ]
    if missing_columns:
        raise KeyError(
            "Resolved Source Registry is missing columns: "
            f"{missing_columns}"
        )

    view = registry[required_columns].copy()
    view["_join_key"] = (
        view["Output File Name"]
        .fillna("")
        .astype(str)
        .str.strip()
        .str.casefold()
    )
    duplicate_keys = view["_join_key"].duplicated(keep=False)
    if duplicate_keys.any():
        values = (
            view.loc[duplicate_keys, "Output File Name"]
            .drop_duplicates()
            .tolist()
        )
        raise ValueError(f"Duplicate Output File Name values: {values}")
    return view.drop(columns=["Output File Name"])


def build_export_dataframe():
    """Read, enrich, normalize, and order the export data."""
    acquisition_log = read_acquisition_log()
    registry_view = build_registry_view()

    acquisition_log["_join_key"] = (
        acquisition_log["output_file"]
        .fillna("")
        .astype(str)
        .str.strip()
        .str.casefold()
    )

    enriched = acquisition_log.merge(
        registry_view,
        on="_join_key",
        how="left",
        validate="many_to_one",
        indicator="_registry_merge",
    )
    enriched[JOIN_STATUS_COLUMN] = enriched["_registry_merge"].map(
        {"both": "MATCHED", "left_only": "UNRESOLVED", "right_only": "UNUSED"}
    ).astype(str)

    unresolved_files = (
        enriched.loc[
            enriched[JOIN_STATUS_COLUMN].eq("UNRESOLVED"),
            "output_file",
        ]
        .dropna()
        .drop_duplicates()
        .tolist()
    )
    if unresolved_files:
        logger.warning(
            "Acquisition Log events without Source Registry metadata: %s",
            unresolved_files,
        )

    enriched = enriched.rename(
        columns={
            "run_id": "Run ID",
            "execution_date": "Execution Date UTC",
            "process_type": "Process Type",
            "output_file": "Output File",
            "storage_location": "Storage Location",
            "status": "Status",
            "records_produced": "Records Produced",
            "notes": "Notes",
            "created_at": "Created At UTC",
        }
    )

    for column in EXPORT_COLUMNS:
        if column not in enriched.columns:
            enriched[column] = None

    export_data = enriched[EXPORT_COLUMNS].copy()
    for column in ("Execution Date UTC", "Created At UTC"):
        export_data[column] = (
            pd.to_datetime(export_data[column], errors="coerce", utc=True)
            .dt.tz_convert(None)
        )
    return export_data


def add_formula_fill(worksheet, column_number, text, fill):
    """Apply a text-equality conditional format to one column."""
    if worksheet.max_row < 2:
        return
    letter = get_column_letter(column_number)
    worksheet.conditional_formatting.add(
        f"{letter}2:{letter}{worksheet.max_row}",
        FormulaRule(formula=[f'${letter}2="{text}"'], fill=fill),
    )


def format_workbook(file_path, row_count):
    """Apply professional formatting to the generated workbook."""
    workbook = load_workbook(file_path)
    worksheet = workbook[SHEET_NAME]
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 30

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = Border(bottom=THIN_GRAY_SIDE)
        worksheet.column_dimensions[get_column_letter(cell.column)].width = (
            COLUMN_WIDTHS.get(str(cell.value), 18)
        )

    header_map = {str(cell.value): cell.column for cell in worksheet[1]}

    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)

    for column_name in ("Execution Date UTC", "Created At UTC"):
        column_number = header_map.get(column_name)
        if column_number is not None:
            for row_number in range(2, worksheet.max_row + 1):
                worksheet.cell(row_number, column_number).number_format = (
                    "yyyy-mm-dd hh:mm:ss"
                )

    records_column = header_map.get("Records Produced")
    if records_column is not None:
        for row_number in range(2, worksheet.max_row + 1):
            worksheet.cell(row_number, records_column).number_format = "#,##0"

    status_column = header_map.get("Status")
    if status_column is not None:
        add_formula_fill(worksheet, status_column, "Success", SUCCESS_FILL)
        add_formula_fill(
            worksheet, status_column, "Partial Success", PARTIAL_FILL
        )
        add_formula_fill(worksheet, status_column, "Failed", FAILED_FILL)
        add_formula_fill(worksheet, status_column, "Planned", PLANNED_FILL)

    match_column = header_map.get(JOIN_STATUS_COLUMN)
    if match_column is not None:
        add_formula_fill(
            worksheet, match_column, "UNRESOLVED", UNRESOLVED_FILL
        )

    if row_count > 0:
        last_column = get_column_letter(worksheet.max_column)
        table = Table(
            displayName=TABLE_NAME,
            ref=f"A1:{last_column}{worksheet.max_row}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(file_path)


def validate_generated_workbook(file_path, expected_rows):
    """Reopen and validate the generated workbook before publication."""
    workbook = load_workbook(file_path, read_only=True, data_only=False)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise RuntimeError(f"Generated workbook is missing sheet: {SHEET_NAME}")
        worksheet = workbook[SHEET_NAME]
        actual_headers = [cell.value for cell in worksheet[1]]
        if actual_headers != EXPORT_COLUMNS:
            raise RuntimeError(
                f"Generated workbook headers are invalid: {actual_headers}"
            )
        actual_rows = max(worksheet.max_row - 1, 0)
        if actual_rows != expected_rows:
            raise RuntimeError(
                f"Generated workbook row count mismatch. "
                f"Expected {expected_rows}, found {actual_rows}."
            )
    finally:
        workbook.close()


def write_excel_atomically(dataframe):
    """Create, validate, and atomically publish the Excel export."""
    ACQUISITION_LOG_FILE.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = None
    try:
        with NamedTemporaryFile(
            mode="wb",
            suffix=".xlsx",
            prefix="acquisition_log_",
            dir=ACQUISITION_LOG_FILE.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)

        dataframe.to_excel(
            temporary_path,
            sheet_name=SHEET_NAME,
            index=False,
            engine="openpyxl",
        )
        format_workbook(temporary_path, len(dataframe))
        validate_generated_workbook(temporary_path, len(dataframe))
        temporary_path.replace(ACQUISITION_LOG_FILE)
    except PermissionError as error:
        raise PermissionError(
            "The Acquisition Log workbook is open or locked. "
            "Close it and rerun the export."
        ) from error
    finally:
        if temporary_path is not None and temporary_path.exists():
            try:
                temporary_path.unlink()
            except OSError as cleanup_error:
                logger.warning(
                    "Unable to remove temporary export %s: %s",
                    temporary_path,
                    cleanup_error,
                )


def export_acquisition_log():
    """Generate the Excel reporting view and return its path."""
    logger.info("Acquisition Log Excel export started")
    try:
        export_data = build_export_dataframe()
        write_excel_atomically(export_data)
        unresolved_count = int(
            export_data[JOIN_STATUS_COLUMN].eq("UNRESOLVED").sum()
        )
        logger.info("Acquisition Log Excel rows exported: %s", len(export_data))
        logger.info(
            "Acquisition Log unresolved registry matches: %s",
            unresolved_count,
        )
        logger.info("Acquisition Log Excel export completed successfully")
        return ACQUISITION_LOG_FILE
    except Exception as error:
        logger.exception("Acquisition Log Excel export failed: %s", error)
        raise


def main():
    print(export_acquisition_log())


if __name__ == "__main__":
    main()
