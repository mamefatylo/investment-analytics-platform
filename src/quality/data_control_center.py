"""
data_control_center.py

Phase 4 - Data Quality Framework

Objectif
--------
Centraliser les contrôles qualité, les indicateurs de qualité,
les anomalies détectées et l'historique des exécutions dans
le Data Control Center.

Produit
-------
audit/data_control_center.xlsx

Feuilles
--------
- Quality Dashboard
- Quality Controls
- Quality Issues
- Quality History
"""

from datetime import datetime
from pathlib import Path

import pandas as pd
import sqlite3

from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.axis import DateAxis
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from src.config import (
    DATA_CONTROL_CENTER_FILE,
    DATABASE_FILE,
)
from src.quality.quality_checks import (
    run_quality_checks,
)
from src.quality.quality_metrics import calculate_quality_metrics
from src.utils.logger import logger

# ============================================================================
# CONFIGURATION
# ============================================================================

DASHBOARD_SHEET = "Quality Dashboard"
CONTROLS_SHEET = "Quality Controls"
ISSUES_SHEET = "Quality Issues"
HISTORY_SHEET = "Quality History"

EXPECTED_SHEETS = [
    DASHBOARD_SHEET,
    CONTROLS_SHEET,
    ISSUES_SHEET,
    HISTORY_SHEET,
]

PASS_THRESHOLD = 95.0
WARNING_THRESHOLD = 80.0

HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="1F4E78",
)

SECTION_FILL = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7",
)

PASS_FILL = PatternFill(
    fill_type="solid",
    fgColor="C6EFCE",
)

WARNING_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFEB9C",
)

FAIL_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFC7CE",
)

NOT_APPLICABLE_FILL = PatternFill(
    fill_type="solid",
    fgColor="D9E1F2",
)

HIGH_FILL = PatternFill(
    fill_type="solid",
    fgColor="F4CCCC",
)

MEDIUM_FILL = PatternFill(
    fill_type="solid",
    fgColor="FCE5CD",
)

LOW_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFF2CC",
)

# ============================================================================
# RUN IDENTIFICATION
# ============================================================================


def generate_run_id(
    execution_datetime: datetime,
) -> str:
    """
    Génère un identifiant unique pour l'exécution qualité.
    """

    return execution_datetime.strftime(
        "DQ-%Y%m%d-%H%M%S-%f"
    )


# ============================================================================
# DATA VALIDATION
# ============================================================================


def validate_quality_results(
    quality_results: dict,
) -> None:
    """
    Vérifie la structure retournée par quality_metrics.py.
    """

    required_keys = {
        "metrics",
        "global_quality_score",
        "global_quality_grade",
        "global_quality_status",
        "records_checked",
        "execution_date",
    }

    missing_keys = (
        required_keys
        - set(quality_results)
    )

    if missing_keys:

        raise KeyError(
            "Missing quality result keys: "
            f"{sorted(missing_keys)}"
        )

    metrics = quality_results["metrics"]

    if not isinstance(
        metrics,
        pd.DataFrame,
    ):

        raise TypeError(
            "quality_results['metrics'] "
            "must be a pandas DataFrame."
        )

    required_metric_columns = {
        "metric",
        "value",
        "status",
        "threshold",
        "records_checked",
        "execution_date",
    }

    missing_columns = (
        required_metric_columns
        - set(metrics.columns)
    )

    if missing_columns:

        raise KeyError(
            "Missing metric columns: "
            f"{sorted(missing_columns)}"
        )

    if metrics.empty:

        raise ValueError(
            "Quality metrics are empty."
        )


def validate_controls(
    controls: pd.DataFrame,
) -> None:
    """
    Vérifie la structure retournée par quality_checks.py.
    """

    required_columns = {
        "control_id",
        "control_name",
        "status",
        "issues",
        "records_checked",
        "execution_date",
        "notes",
    }

    missing_columns = (
        required_columns
        - set(controls.columns)
    )

    if missing_columns:

        raise KeyError(
            "Missing control columns: "
            f"{sorted(missing_columns)}"
        )

    if controls.empty:

        raise ValueError(
            "Quality controls are empty."
        )


# ============================================================================
# DASHBOARD
# ============================================================================


def build_dashboard_summary(
    quality_results: dict,
    run_id: str,
    execution_date: str,
) -> pd.DataFrame:
    """
    Construit la synthèse exécutive du Data Control Center.
    """

    return pd.DataFrame(
        [
            {
                "run_id": run_id,
                "execution_date": execution_date,
                "records_checked":
                    quality_results[
                        "records_checked"
                    ],
                "global_quality_score":
                    quality_results[
                        "global_quality_score"
                    ],
                "global_quality_grade":
                    quality_results[
                        "global_quality_grade"
                    ],
                "global_status":
                    quality_results[
                        "global_quality_status"
                    ],
            }
        ]
    )


def build_dashboard_metrics(
    quality_results: dict,
) -> pd.DataFrame:
    """
    Construit la table des indicateurs affichée dans le dashboard.
    """

    metrics = quality_results[
        "metrics"
    ].copy()

    metrics = metrics[
        [
            "metric",
            "value",
            "status",
            "threshold",
        ]
    ]

    return metrics


# ============================================================================
# QUALITY ISSUES
# ============================================================================


def determine_severity(
    issue_count: int,
    records_checked: int,
) -> str:
    """
    Détermine la sévérité d'une anomalie selon son taux d'incidence.
    """

    if issue_count <= 0:
        return "NONE"

    if records_checked <= 0:
        return "HIGH"

    issue_rate = (
        issue_count
        / records_checked
    ) * 100

    if issue_rate >= 10:
        return "HIGH"

    if issue_rate >= 5:
        return "MEDIUM"

    return "LOW"


def build_quality_issues(
    controls: pd.DataFrame,
    run_id: str,
    execution_date: str,
) -> pd.DataFrame:
    """
    Construit le registre des contrôles présentant des anomalies.
    """

    issue_controls = controls[
        controls["issues"] > 0
    ].copy()

    issue_columns = [
        "issue_id",
        "run_id",
        "control_id",
        "control_name",
        "severity",
        "issue_count",
        "records_checked",
        "issue_rate",
        "detection_date",
        "notes",
    ]

    if issue_controls.empty:

        return pd.DataFrame(
            columns=issue_columns
        )

    issue_controls = (
        issue_controls
        .reset_index(drop=True)
    )

    issue_records = []

    for index, row in (
        issue_controls.iterrows()
    ):

        issue_count = int(
            row["issues"]
        )

        records_checked = int(
            row["records_checked"]
        )

        issue_rate = (
            (
                issue_count
                / records_checked
            ) * 100
            if records_checked > 0
            else 100.0
        )

        issue_records.append(
            {
                "issue_id":
                    f"{run_id}-ISS-{index + 1:03d}",
                "run_id":
                    run_id,
                "control_id":
                    row["control_id"],
                "control_name":
                    row["control_name"],
                "severity":
                    determine_severity(
                        issue_count,
                        records_checked,
                    ),
                "issue_count":
                    issue_count,
                "records_checked":
                    records_checked,
                "issue_rate":
                    round(
                        issue_rate,
                        2,
                    ),
                "detection_date":
                    execution_date,
                "notes":
                    row["notes"],
            }
        )

    return pd.DataFrame(
        issue_records,
        columns=issue_columns,
    )


# ============================================================================
# QUALITY HISTORY
# ============================================================================


def load_history() -> pd.DataFrame:
    """
    Charge l'historique existant du Data Control Center.
    """

    history_columns = [
        "run_id",
        "execution_date",
        "records_checked",
        "availability_rate",
        "completeness_rate",
        "uniqueness_rate",
        "freshness_rate",
        "chronology_rate",
        "global_quality_score",
        "global_quality_grade",
        "global_status",
    ]

    if not DATA_CONTROL_CENTER_FILE.exists():

        return pd.DataFrame(
            columns=history_columns
        )

    try:

        history = pd.read_excel(
            DATA_CONTROL_CENTER_FILE,
            sheet_name=HISTORY_SHEET,
            engine="openpyxl",
        )

    except ValueError:

        logger.warning(
            "Quality History sheet not found. "
            "A new history will be created."
        )

        return pd.DataFrame(
            columns=history_columns
        )

    except PermissionError as error:

        raise PermissionError(
            "The Data Control Center is currently "
            "open or locked. Close the workbook "
            "and rerun the process."
        ) from error

    except Exception as error:

        logger.exception(
            "Unable to read existing quality history."
        )

        raise RuntimeError(
            "Unable to read existing "
            "Quality History."
        ) from error

    for column in history_columns:

        if column not in history.columns:
            history[column] = None

    return history[
        history_columns
    ]


def get_metric_value(
    metrics: pd.DataFrame,
    metric_name: str,
) -> float:
    """
    Récupère la valeur d'un indicateur par son nom.
    """

    metric_rows = metrics[
        metrics["metric"] == metric_name
    ]

    if metric_rows.empty:

        raise KeyError(
            f"Metric not found: {metric_name}"
        )

    return float(
        metric_rows.iloc[0]["value"]
    )


def update_history(
    quality_results: dict,
    run_id: str,
    execution_date: str,
) -> pd.DataFrame:
    """
    Ajoute l'exécution courante à l'historique qualité.
    """

    history = load_history()

    metrics = quality_results[
        "metrics"
    ]

    current_run = {
        "run_id":
            run_id,

        "execution_date":
            execution_date,

        "records_checked":
            quality_results[
                "records_checked"
            ],

        "availability_rate":
            get_metric_value(
                metrics,
                "Availability Rate",
            ),

        "completeness_rate":
            get_metric_value(
                metrics,
                "Completeness Rate",
            ),

        "uniqueness_rate":
            get_metric_value(
                metrics,
                "Uniqueness Rate",
            ),

        "freshness_rate":
            get_metric_value(
                metrics,
                "Freshness Rate",
            ),

        "chronology_rate":
            get_metric_value(
                metrics,
                "Chronology Rate",
            ),

        "global_quality_score":
            quality_results[
                "global_quality_score"
            ],

        "global_quality_grade":
            quality_results[
                "global_quality_grade"
            ],

        "global_status":
            quality_results[
                "global_quality_status"
            ],
    }

    updated_history = pd.concat(
        [
            history,
            pd.DataFrame(
                [current_run]
            ),
        ],
        ignore_index=True,
    )

    updated_history[
        "execution_date"
    ] = pd.to_datetime(
        updated_history[
            "execution_date"
        ],
        errors="coerce",
    )

    updated_history = (
        updated_history
        .sort_values(
            by="execution_date",
            kind="stable",
        )
        .reset_index(drop=True)
    )

    return updated_history


# ============================================================================
# FILE ACCESS
# ============================================================================


def validate_output_file_access() -> None:
    """
    Vérifie que le fichier Excel peut être remplacé.
    """

    if not DATA_CONTROL_CENTER_FILE.exists():
        return

    try:

        with DATA_CONTROL_CENTER_FILE.open(
            mode="a",
        ):
            pass

    except PermissionError as error:

        raise PermissionError(
            "The Data Control Center is currently "
            "open or locked. Close "
            "audit/data_control_center.xlsx "
            "and rerun the process."
        ) from error


# ============================================================================
# EXCEL FORMATTING
# ============================================================================


def style_header_row(
    worksheet: Worksheet,
    row_number: int,
    max_column: int,
) -> None:
    """
    Applique le style standard à une ligne d'en-tête.
    """

    for cell in worksheet[
        row_number
    ][:max_column]:

        cell.fill = HEADER_FILL

        cell.font = Font(
            color="FFFFFF",
            bold=True,
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )


def style_section_title(
    worksheet: Worksheet,
    cell_reference: str,
) -> None:
    """
    Applique le style standard à un titre de section.
    """

    cell = worksheet[
        cell_reference
    ]

    cell.fill = SECTION_FILL

    cell.font = Font(
        bold=True,
        color="1F1F1F",
        size=12,
    )


def set_column_widths(
    worksheet: Worksheet,
    minimum_width: int = 12,
    maximum_width: int = 45,
) -> None:
    """
    Ajuste automatiquement la largeur des colonnes.
    """

    for column_cells in worksheet.columns:

        column_letter = get_column_letter(
            column_cells[0].column
        )

        maximum_length = 0

        for cell in column_cells:

            if cell.value is None:
                continue

            cell_length = len(
                str(cell.value)
            )

            maximum_length = max(
                maximum_length,
                cell_length,
            )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(
                maximum_length + 2,
                minimum_width,
            ),
            maximum_width,
        )


def apply_status_formatting(
    worksheet: Worksheet,
    status_column: int,
    first_data_row: int,
    last_data_row: int,
) -> None:
    """
    Applique une couleur conditionnelle aux statuts.
    """

    column_letter = get_column_letter(
        status_column
    )

    cell_range = (
        f"{column_letter}{first_data_row}:"
        f"{column_letter}{last_data_row}"
    )

    worksheet.conditional_formatting.add(
        cell_range,
        CellIsRule(
            operator="equal",
            formula=['"PASS"'],
            fill=PASS_FILL,
        ),
    )

    worksheet.conditional_formatting.add(
        cell_range,
        CellIsRule(
            operator="equal",
            formula=['"WARNING"'],
            fill=WARNING_FILL,
        ),
    )

    worksheet.conditional_formatting.add(
        cell_range,
        CellIsRule(
            operator="equal",
            formula=['"FAIL"'],
            fill=FAIL_FILL,
        ),
    )

    worksheet.conditional_formatting.add(
        cell_range,
        CellIsRule(
            operator="equal",
            formula=['"NOT_APPLICABLE"'],
            fill=NOT_APPLICABLE_FILL,
        ),
    )


def format_dashboard_sheet(
    worksheet: Worksheet,
    summary_columns: int,
    metrics_header_row: int,
    metrics_rows: int,
) -> None:
    """
    Met en forme la feuille Quality Dashboard.
    """

    worksheet.freeze_panes = "A5"

    style_section_title(
        worksheet,
        "A1",
    )

    style_header_row(
        worksheet,
        row_number=2,
        max_column=summary_columns,
    )

    style_section_title(
        worksheet,
        "A4",
    )

    style_header_row(
        worksheet,
        row_number=metrics_header_row,
        max_column=4,
    )

    worksheet[
        "D3"
    ].number_format = "0.00"

    for row in range(
        metrics_header_row + 1,
        metrics_header_row
        + metrics_rows
        + 1,
    ):

        worksheet.cell(
            row=row,
            column=2,
        ).number_format = "0.00"

        worksheet.cell(
            row=row,
            column=4,
        ).number_format = "0.00"

    apply_status_formatting(
        worksheet,
        status_column=6,
        first_data_row=3,
        last_data_row=3,
    )

    apply_status_formatting(
        worksheet,
        status_column=3,
        first_data_row=metrics_header_row + 1,
        last_data_row=(
            metrics_header_row
            + metrics_rows
        ),
    )

    set_column_widths(
        worksheet
    )


def format_controls_sheet(
    worksheet: Worksheet,
) -> None:
    """
    Met en forme la feuille Quality Controls.
    """

    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    style_header_row(
        worksheet,
        row_number=1,
        max_column=worksheet.max_column,
    )

    headers = {
        cell.value: cell.column
        for cell in worksheet[1]
    }

    status_column = headers.get(
        "status"
    )

    if status_column:

        apply_status_formatting(
            worksheet,
            status_column=status_column,
            first_data_row=2,
            last_data_row=worksheet.max_row,
        )

    set_column_widths(
        worksheet
    )


def format_issues_sheet(
    worksheet: Worksheet,
) -> None:
    """
    Met en forme la feuille Quality Issues.
    """

    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    style_header_row(
        worksheet,
        row_number=1,
        max_column=worksheet.max_column,
    )

    headers = {
        cell.value: cell.column
        for cell in worksheet[1]
    }

    severity_column = headers.get(
        "severity"
    )

    if severity_column:

        column_letter = get_column_letter(
            severity_column
        )

        cell_range = (
            f"{column_letter}2:"
            f"{column_letter}{worksheet.max_row}"
        )

        worksheet.conditional_formatting.add(
            cell_range,
            CellIsRule(
                operator="equal",
                formula=['"HIGH"'],
                fill=HIGH_FILL,
            ),
        )

        worksheet.conditional_formatting.add(
            cell_range,
            CellIsRule(
                operator="equal",
                formula=['"MEDIUM"'],
                fill=MEDIUM_FILL,
            ),
        )

        worksheet.conditional_formatting.add(
            cell_range,
            CellIsRule(
                operator="equal",
                formula=['"LOW"'],
                fill=LOW_FILL,
            ),
        )

    set_column_widths(
        worksheet
    )


def format_issues_sheet(
    worksheet: Worksheet,
) -> None:
    """
    Met en forme la feuille Quality Issues.

    La mise en forme conditionnelle des sévérités
    est appliquée uniquement lorsque la feuille
    contient au moins une anomalie.
    """

    worksheet.freeze_panes = "A2"

    style_header_row(
        worksheet,
        row_number=1,
        max_column=worksheet.max_column,
    )

    set_column_widths(
        worksheet
    )

    # La feuille ne contient que les en-têtes.
    # Aucune plage de données ne doit être formatée.
    if worksheet.max_row < 2:
        return

    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    headers = {
        cell.value: cell.column
        for cell in worksheet[1]
    }

    severity_column = headers.get(
        "severity"
    )

    if severity_column is None:
        return

    column_letter = get_column_letter(
        severity_column
    )

    cell_range = (
        f"{column_letter}2:"
        f"{column_letter}{worksheet.max_row}"
    )

    worksheet.conditional_formatting.add(
        cell_range,
        CellIsRule(
            operator="equal",
            formula=['"HIGH"'],
            fill=HIGH_FILL,
        ),
    )

    worksheet.conditional_formatting.add(
        cell_range,
        CellIsRule(
            operator="equal",
            formula=['"MEDIUM"'],
            fill=MEDIUM_FILL,
        ),
    )

    worksheet.conditional_formatting.add(
        cell_range,
        CellIsRule(
            operator="equal",
            formula=['"LOW"'],
            fill=LOW_FILL,
        ),
    )

def format_history_sheet(
    worksheet: Worksheet,
) -> None:
    """
    Met en forme la feuille Quality History.

    La fonction applique :
    - un style aux en-têtes ;
    - un filtre automatique ;
    - le format chronologique aux dates ;
    - le format numérique aux scores ;
    - une mise en forme conditionnelle aux statuts ;
    - un ajustement automatique des colonnes.
    """

    worksheet.freeze_panes = "A2"

    style_header_row(
        worksheet,
        row_number=1,
        max_column=worksheet.max_column,
    )

    set_column_widths(
        worksheet
    )

    # Aucun enregistrement historique.
    if worksheet.max_row < 2:
        return

    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    headers = {
        cell.value: cell.column
        for cell in worksheet[1]
    }

    execution_date_column = headers.get(
        "execution_date"
    )

    global_score_column = headers.get(
        "global_quality_score"
    )

    global_status_column = headers.get(
        "global_status"
    )

    if execution_date_column is not None:

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):

            worksheet.cell(
                row=row_number,
                column=execution_date_column,
            ).number_format = (
                "yyyy-mm-dd hh:mm:ss"
            )

    if global_score_column is not None:

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):

            worksheet.cell(
                row=row_number,
                column=global_score_column,
            ).number_format = "0.00"

    if global_status_column is not None:

        apply_status_formatting(
            worksheet,
            status_column=global_status_column,
            first_data_row=2,
            last_data_row=worksheet.max_row,
        )
# ============================================================================
# EXCEL CHARTS
# ============================================================================


def create_kpi_chart(
    worksheet: Worksheet,
    header_row: int,
    number_of_metrics: int,
) -> None:
    """
    Crée le graphique des indicateurs qualité.
    """

    if number_of_metrics <= 0:
        return

    chart = BarChart()

    chart.type = "col"

    chart.style = 10

    chart.title = (
        "Quality KPI Overview"
    )

    chart.y_axis.title = (
        "Score (%)"
    )

    chart.x_axis.title = (
        "Quality Metric"
    )

    chart.y_axis.scaling.min = 0

    chart.y_axis.scaling.max = 100

    chart.height = 9

    chart.width = 18

    data = Reference(
        worksheet,
        min_col=2,
        min_row=header_row,
        max_row=(
            header_row
            + number_of_metrics
        ),
    )

    categories = Reference(
        worksheet,
        min_col=1,
        min_row=header_row + 1,
        max_row=(
            header_row
            + number_of_metrics
        ),
    )

    chart.add_data(
        data,
        titles_from_data=True,
    )

    chart.set_categories(
        categories
    )

    chart.legend = None

    chart.dataLabels = DataLabelList()

    chart.dataLabels.showVal = True

    worksheet.add_chart(
        chart,
        "H2",
    )


def create_history_chart(
    worksheet: Worksheet,
) -> None:
    """
    Crée le graphique chronologique du score global.
    """

    if worksheet.max_row < 3:
        return

    headers = {
        cell.value: cell.column
        for cell in worksheet[1]
    }

    score_column = headers.get(
        "global_quality_score"
    )

    date_column = headers.get(
        "execution_date"
    )

    if (
        score_column is None
        or date_column is None
    ):
        return

    chart = LineChart()

    chart.style = 13

    chart.title = (
        "Global Quality Score Trend"
    )

    chart.y_axis.title = (
        "Global Quality Score (%)"
    )

    chart.x_axis = DateAxis(
        crossAx=100
    )

    chart.x_axis.title = (
        "Execution Date"
    )

    chart.x_axis.number_format = (
        "yyyy-mm-dd hh:mm"
    )

    chart.x_axis.majorTimeUnit = (
        "days"
    )

    chart.y_axis.scaling.min = 0

    chart.y_axis.scaling.max = 100

    chart.height = 10

    chart.width = 19

    data = Reference(
        worksheet,
        min_col=score_column,
        min_row=1,
        max_row=worksheet.max_row,
    )

    dates = Reference(
        worksheet,
        min_col=date_column,
        min_row=2,
        max_row=worksheet.max_row,
    )

    chart.add_data(
        data,
        titles_from_data=True,
    )

    chart.set_categories(
        dates
    )

    chart.legend = None

    if chart.series:

        series = chart.series[0]

        series.graphicalProperties.line.solidFill = (
            "1F4E78"
        )

        series.graphicalProperties.line.width = (
            28575
        )

        series.marker.symbol = "circle"

        series.marker.size = 7

    worksheet.add_chart(
        chart,
        "N2",
    )


# ============================================================================
# WORKBOOK GENERATION
# ============================================================================


def write_workbook(
    dashboard_summary: pd.DataFrame,
    dashboard_metrics: pd.DataFrame,
    controls: pd.DataFrame,
    issues: pd.DataFrame,
    history: pd.DataFrame,
) -> None:
    """
    Écrit les quatre feuilles du Data Control Center.
    """

    validate_output_file_access()

    DATA_CONTROL_CENTER_FILE.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        DATA_CONTROL_CENTER_FILE,
        engine="openpyxl",
        mode="w",
        datetime_format=(
            "yyyy-mm-dd hh:mm:ss"
        ),
    ) as writer:

        dashboard_summary.to_excel(
            writer,
            sheet_name=DASHBOARD_SHEET,
            index=False,
            startrow=1,
        )

        dashboard_metrics.to_excel(
            writer,
            sheet_name=DASHBOARD_SHEET,
            index=False,
            startrow=4,
        )

        controls.to_excel(
            writer,
            sheet_name=CONTROLS_SHEET,
            index=False,
        )

        issues.to_excel(
            writer,
            sheet_name=ISSUES_SHEET,
            index=False,
        )

        history.to_excel(
            writer,
            sheet_name=HISTORY_SHEET,
            index=False,
        )

        dashboard_worksheet = (
            writer.book[
                DASHBOARD_SHEET
            ]
        )

        dashboard_worksheet[
            "A1"
        ] = "Executive Summary"

        dashboard_worksheet[
            "A4"
        ] = "Quality Metrics"


def finalize_workbook(
    dashboard_summary: pd.DataFrame,
    dashboard_metrics: pd.DataFrame,
) -> None:
    """
    Applique la mise en forme et ajoute les graphiques.
    """

    workbook = load_workbook(
        DATA_CONTROL_CENTER_FILE
    )

    if workbook.sheetnames != EXPECTED_SHEETS:

        raise ValueError(
            "Unexpected workbook structure. "
            f"Expected {EXPECTED_SHEETS}, "
            f"received {workbook.sheetnames}."
        )

    dashboard_worksheet = workbook[
        DASHBOARD_SHEET
    ]

    controls_worksheet = workbook[
        CONTROLS_SHEET
    ]

    issues_worksheet = workbook[
        ISSUES_SHEET
    ]

    history_worksheet = workbook[
        HISTORY_SHEET
    ]

    format_dashboard_sheet(
        dashboard_worksheet,
        summary_columns=len(
            dashboard_summary.columns
        ),
        metrics_header_row=5,
        metrics_rows=len(
            dashboard_metrics
        ),
    )

    format_controls_sheet(
        controls_worksheet
    )

    format_issues_sheet(
        issues_worksheet
    )

    format_history_sheet(
        history_worksheet
    )

    create_kpi_chart(
        dashboard_worksheet,
        header_row=5,
        number_of_metrics=len(
            dashboard_metrics
        ),
    )

    create_history_chart(
        history_worksheet
    )

    workbook.save(
        DATA_CONTROL_CENTER_FILE
    )


# ============================================================================
# DATA CONTROL CENTER
# ============================================================================

def build_data_control_center() -> Path:
    """
    Génère le Data Control Center complet.
    """

    logger.info(
        "Data Control Center generation started"
    )

    execution_datetime = datetime.now()

    execution_date = (
        execution_datetime.strftime(
            "%Y-%m-%d %H:%M:%S"
        )
    )

    run_id = generate_run_id(
        execution_datetime
    )

    try:

        controls = (
            run_quality_checks()
        )

        validate_controls(
            controls
        )

        connection = sqlite3.connect(
            DATABASE_FILE
        )

        try:

            dataframe = pd.read_sql(
                """
                SELECT *
                FROM securities_master
                """,
                connection,
            )

        finally:

            connection.close()

        quality_results = (
            calculate_quality_metrics(
                dataframe=dataframe,
                chronology_enabled=False,
            )
        )

        validate_quality_results(
            quality_results
        )

        dashboard_summary = (
            build_dashboard_summary(
                quality_results,
                run_id,
                execution_date,
            )
        )

        dashboard_metrics = (
            build_dashboard_metrics(
                quality_results
            )
        )

        issues = (
            build_quality_issues(
                controls,
                run_id,
                execution_date,
            )
        )

        history = (
            update_history(
                quality_results,
                run_id,
                execution_date,
            )
        )

        write_workbook(
            dashboard_summary,
            dashboard_metrics,
            controls,
            issues,
            history,
        )

        finalize_workbook(
            dashboard_summary,
            dashboard_metrics,
        )

        logger.info(
            f"Quality run ID: {run_id}"
        )

        logger.info(
            f"Records checked: "
            f"{quality_results['records_checked']:,}"
        )

        logger.info(
            f"Controls executed: "
            f"{len(controls)}"
        )

        logger.info(
            f"Issue categories detected: "
            f"{len(issues)}"
        )

        logger.info(
            f"Global Quality Score: "
            f"{quality_results['global_quality_score']:.2f}"
        )

        logger.info(
            f"Global Quality Grade: "
            f"{quality_results['global_quality_grade']}"
        )

        logger.info(
            f"Global Quality Status: "
            f"{quality_results['global_quality_status']}"
        )

        logger.info(
            f"History records: "
            f"{len(history)}"
        )

        logger.info(
            f"Output file: "
            f"{DATA_CONTROL_CENTER_FILE}"
        )

        logger.info(
            "Data Control Center generation completed successfully"
        )

        return DATA_CONTROL_CENTER_FILE

    except Exception as error:

        logger.exception(
            "Data Control Center generation failed: "
            f"{error}"
        )

        raise

# ============================================================================
# MAIN
# ============================================================================


def main() -> None:
    """
    Point d'entrée du script.
    """

    build_data_control_center()


if __name__ == "__main__":

    main()